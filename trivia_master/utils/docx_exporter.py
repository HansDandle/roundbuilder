"""Export trivia rounds to DOCX format"""
from typing import List, Dict, Any
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


class DocxExporter:
    """Export trivia rounds to DOCX files"""
    
    def __init__(self):
        self.output_dir = Path('./output')
        self.output_dir.mkdir(exist_ok=True)
    
    def export_round(self, 
                    round_data: Dict[str, Any],
                    filename: str = None,
                    include_answers: bool = True) -> str:
        """Export a single round to DOCX"""
        if not filename:
            filename = f"{round_data.get('name', 'Round')}.docx"
        
        doc = Document()
        
        # Add title
        title = doc.add_heading(round_data.get('name', 'Trivia Round'), level=1)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Add metadata
        if 'question_count' in round_data:
            doc.add_paragraph(f"Total Questions: {round_data['question_count']}")
        
        if 'categories' in round_data and round_data['categories']:
            doc.add_paragraph(f"Categories: {', '.join(round_data['categories'])}")
        
        # Add questions
        questions = round_data.get('questions', [])
        for i, question in enumerate(questions, 1):
            # Question
            q_para = doc.add_paragraph(f"{i}. {question['question']}", style='List Number')
            q_run = q_para.runs[0]
            q_run.bold = True
            
            # Answer (if included)
            if include_answers:
                answer_text = question['answer']
                if question.get('category'):
                    answer_text = f"Answer: {answer_text} (Category: {question['category']})"
                else:
                    answer_text = f"Answer: {answer_text}"
                
                ans_para = doc.add_paragraph(answer_text)
                ans_run = ans_para.runs[0]
                ans_run.italic = True
            else:
                # Leave space for written answers
                doc.add_paragraph('_' * 80)
            
            # Add spacing
            doc.add_paragraph()
        
        # Save file
        filepath = self.output_dir / filename
        doc.save(str(filepath))
        return str(filepath)
    
    def export_multiple_rounds(self,
                              rounds: List[Dict[str, Any]],
                              filename: str = None) -> str:
        """Export multiple rounds to a single DOCX file"""
        if not filename:
            filename = "Multiple_Rounds.docx"
        
        doc = Document()
        
        # Add main title
        title = doc.add_heading("Trivia Rounds Collection", level=1)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_paragraph(f"Total Rounds: {len(rounds)}")
        doc.add_page_break()
        
        # Add each round
        for round_idx, round_data in enumerate(rounds, 1):
            # Round title
            round_title = doc.add_heading(f"Round {round_idx}: {round_data.get('name', 'Round')}", level=2)
            
            if 'question_count' in round_data:
                doc.add_paragraph(f"Questions: {round_data['question_count']}")
            
            # Questions
            questions = round_data.get('questions', [])
            for i, question in enumerate(questions, 1):
                q_para = doc.add_paragraph(f"{i}. {question['question']}", style='List Number')
                q_run = q_para.runs[0]
                q_run.bold = True
                
                answer_text = question['answer']
                if question.get('category'):
                    answer_text = f"Answer: {answer_text} (Category: {question['category']})"
                else:
                    answer_text = f"Answer: {answer_text}"
                
                ans_para = doc.add_paragraph(answer_text)
                ans_run = ans_para.runs[0]
                ans_run.italic = True
                
                doc.add_paragraph()
            
            if round_idx < len(rounds):
                doc.add_page_break()
        
        filepath = self.output_dir / filename
        doc.save(str(filepath))
        return str(filepath)
    
    def set_output_directory(self, directory: str):
        """Set custom output directory"""
        self.output_dir = Path(directory)
        self.output_dir.mkdir(parents=True, exist_ok=True)
